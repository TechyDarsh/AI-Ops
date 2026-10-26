import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install("python-docx")
    import docx

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

doc = Document()

snippets = [
    {
        "title": "1. Arduino Firmware (Sensors & Serial Comm)",
        "code": """#include <DHT.h>
#define DHTPIN 2
#define DHTTYPE DHT22
DHT dht(DHTPIN, DHTTYPE);
#define MQ2_PIN A0
#define MQ9_PIN A1
void setup() {
  Serial.begin(9600);
  dht.begin();
}
void loop() {
  delay(2000); 
  float h = dht.readHumidity();
  float t = dht.readTemperature();
  if (isnan(h) || isnan(t)) { h = 0.0; t = 0.0; }
  int mq2_raw = analogRead(MQ2_PIN);
  float mq2_voltage = mq2_raw * (5.0 / 1023.0);
  int mq9_raw = analogRead(MQ9_PIN);
  float mq9_voltage = mq9_raw * (5.0 / 1023.0);
  
  Serial.print("{");
  Serial.print("\"deviceId\":\"ARDUINO-001\",");
  Serial.print("\"source\":\"REAL\",");
  Serial.print("\"dht22\":{\"temperature\":");
  Serial.print(t, 2);
  Serial.print(",\"humidity\":");
  Serial.print(h, 2);
  Serial.print("},");
  Serial.print("\"mq2\":{\"raw\":");
  Serial.print(mq2_raw);
  Serial.print(",\"voltage\":");
  Serial.print(mq2_voltage, 2);
  Serial.print("},");
  Serial.print("\"mq9\":{\"raw\":");
  Serial.print(mq9_raw);
  Serial.print(",\"voltage\":");
  Serial.print(mq9_voltage, 2);
  Serial.println("}}");
}"""
    },
    {
        "title": "2. Python Cloud Server (WebSocket Relay)",
        "code": """import asyncio
import websockets

clients = set()
latest_data = '{"status": "AWAITING_DATA"}'

async def relay_handler(websocket, path):
    global latest_data
    clients.add(websocket)
    try:
        await websocket.send(latest_data)
    except:
        pass
    try:
        async for message in websocket:
            latest_data = message
            for client in list(clients):
                if client != websocket:
                    try:
                        await client.send(message)
                    except websockets.exceptions.ConnectionClosed:
                        clients.remove(client)
    except websockets.exceptions.ConnectionClosed:
        pass
    except Exception as e:
        print(f"Error: {e}")
    finally:
        clients.remove(websocket)

async def main():
    port = 8081
    async with websockets.serve(relay_handler, "0.0.0.0", port):
        await asyncio.Future()

if __name__ == "__main__":
    asyncio.run(main())"""
    },
    {
        "title": "3. Gateway Script (WebSocket Broadcast)",
        "code": """import asyncio
import websockets
import json

connected_clients = set()
latest_telemetry = {"deviceId": "GW-001", "status": "ONLINE"}

async def ws_handler(websocket, path):
    connected_clients.add(websocket)
    try:
        while True:
            if latest_telemetry:
                await websocket.send(json.dumps(latest_telemetry))
            await asyncio.sleep(1)
    except websockets.exceptions.ConnectionClosed:
        pass
    finally:
        connected_clients.remove(websocket)

async def main():
    port = 8081
    async with websockets.serve(ws_handler, "0.0.0.0", port):
        await asyncio.Future()

if __name__ == "__main__":
    asyncio.run(main())"""
    },
    {
        "title": "4. API Service (Frontend API Layer)",
        "code": """const BASE_URL = 'http://localhost:8080/api';
function getToken() { return localStorage.getItem('aiops_token'); }
function getHeaders() {
  const token = getToken();
  const headers = { 'Content-Type': 'application/json' };
  if (token) headers['Authorization'] = `Bearer ${token}`;
  return headers;
}
async function request(method, path, body = null) {
  const options = { method, headers: getHeaders() };
  if (body) options.body = JSON.stringify(body);
  try {
    const res = await fetch(`${BASE_URL}${path}`, options);
    if (res.status === 401) {
      if (localStorage.getItem('aiops_token')) {
        localStorage.removeItem('aiops_token');
        window.dispatchEvent(new Event('auth-expired'));
      }
      return null;
    }
    const text = await res.text();
    if (!text || !text.trim()) return null;
    try { return JSON.parse(text); }
    catch { return null; }
  } catch (err) {
    throw err;
  }
}
export const authApi = {
  login: (email, password) => request('POST', '/auth/login', { email, password }),
  logout: () => request('POST', '/auth/logout'),
  me: () => request('GET', '/auth/me'),
};"""
    },
    {
        "title": "5. Java JWT Utils (Security)",
        "code": """package com.aiops.security;

import io.jsonwebtoken.*;
import io.jsonwebtoken.io.Decoders;
import io.jsonwebtoken.security.Keys;
import org.springframework.beans.factory.annotation.Value;
import org.springframework.security.core.userdetails.UserDetails;
import org.springframework.stereotype.Component;
import javax.crypto.SecretKey;
import java.util.Date;
import java.util.HashMap;
import java.util.Map;

@Component
public class JwtUtils {
    @Value("${aiops.jwt.secret}")
    private String jwtSecret;
    @Value("${aiops.jwt.expiration-ms}")
    private int jwtExpirationMs;

    private SecretKey getSigningKey() {
        byte[] keyBytes = Decoders.BASE64.decode(
            java.util.Base64.getEncoder().encodeToString(jwtSecret.getBytes())
        );
        return Keys.hmacShaKeyFor(keyBytes);
    }
    
    public String generateToken(UserDetails userDetails, String role) {
        Map<String, Object> claims = new HashMap<>();
        claims.put("role", role);
        return Jwts.builder()
                .claims(claims)
                .subject(userDetails.getUsername())
                .issuedAt(new Date())
                .expiration(new Date(System.currentTimeMillis() + jwtExpirationMs))
                .signWith(getSigningKey())
                .compact();
    }
}"""
    }
]

for i, snippet in enumerate(snippets):
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(snippet["title"])
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(14)
    
    p_code = doc.add_paragraph()
    r_code = p_code.add_run(snippet["code"])
    r_code.font.name = 'Times New Roman'
    r_code.font.size = Pt(12)
    
    if i < len(snippets) - 1:
        # We want to add a page break. Let's do it on a new paragraph to be safe.
        p_break = doc.add_paragraph()
        p_break.add_run().add_break(WD_BREAK.PAGE)

output_dir = r"C:\Users\darsh\OneDrive\Desktop\desk\Project\word"
if not os.path.exists(output_dir):
    os.makedirs(output_dir)

doc.save(os.path.join(output_dir, "Sample_Coding_Documentation.docx"))
print("Word document generated successfully!")
